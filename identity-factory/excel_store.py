"""Append-only Excel workbook of generated Copley identities, one sheet per lifecycle."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from user_payload import INITIALS, LIFECYCLES, MY_RUSH_JOBS_COLUMNS

DATA_DIR = Path(__file__).resolve().parent / "data"
EXCEL_PATH = DATA_DIR / "Copley_test_identities.xlsx"

META_COLUMNS = ("Created_At", "Lifecycle")
HEADERS = META_COLUMNS + MY_RUSH_JOBS_COLUMNS


def _style_header(ws: Worksheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col, title)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
    widths = {"Created_At": 20, "Lifecycle": 14, "Stage_Key": 28, "Display_Name": 24, "Work_Email": 32}
    for col, title in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(title, 18)


def _workbook() -> tuple:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if EXCEL_PATH.exists():
        wb = load_workbook(EXCEL_PATH)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)
    for name in LIFECYCLES:
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name)
            _style_header(ws)
        else:
            ws = wb[name]
            if ws.cell(1, 1).value != "Created_At":
                ws.insert_rows(1)
                _style_header(ws)
    return wb


def used_stage_keys() -> set[str]:
    if not EXCEL_PATH.exists():
        return set()
    keys: set[str] = set()
    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    try:
        idx = HEADERS.index("Stage_Key")
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=2, min_col=idx + 1, max_col=idx + 1, values_only=True):
                if row[0]:
                    keys.add(str(row[0]))
    finally:
        wb.close()
    return keys


def used_first_last() -> set[tuple[str, str]]:
    """Original first+last pairs already stored (AA suffix stripped from Given_Name)."""
    if not EXCEL_PATH.exists():
        return set()
    used: set[tuple[str, str]] = set()
    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    try:
        given_i = HEADERS.index("Given_Name")
        family_i = HEADERS.index("Family_Name")
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=2, min_col=1, max_col=len(HEADERS), values_only=True):
                given = str(row[given_i] or "").strip()
                family = str(row[family_i] or "").strip()
                if not given or not family:
                    continue
                first = given[: -len(INITIALS)] if given.endswith(INITIALS) and len(given) > len(INITIALS) else given
                used.add((first, family))
    finally:
        wb.close()
    return used


def append_identity(lifecycle: str, row: dict) -> Path:
    if lifecycle not in LIFECYCLES:
        raise SystemExit(f"Unknown lifecycle sheet: {lifecycle}")
    wb = _workbook()
    ws = wb[lifecycle]
    stage = row.get("Stage_Key")
    for existing in ws.iter_rows(min_row=2, min_col=HEADERS.index("Stage_Key") + 1, max_col=HEADERS.index("Stage_Key") + 1, values_only=True):
        if existing[0] == stage:
            wb.save(EXCEL_PATH)
            return EXCEL_PATH
    values = [datetime.now().strftime("%Y-%m-%d %H:%M:%S"), lifecycle]
    values.extend("" if row.get(col) is None else row.get(col) for col in MY_RUSH_JOBS_COLUMNS)
    ws.append(values)
    try:
        wb.save(EXCEL_PATH)
    except PermissionError:
        raise SystemExit(
            f"Could not write {EXCEL_PATH}. Close the Excel file if it is open, then run again."
        ) from None
    return EXCEL_PATH
